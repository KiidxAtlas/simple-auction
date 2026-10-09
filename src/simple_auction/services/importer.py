"""Import lots from old auction spreadsheets (.xlsx or .csv) of any layout.

The header row is found automatically, columns are matched to lot fields by
their names, and the caller can override any match before building lots.
"""

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from simple_auction.constants import AUCTION_STEP, TITLE_MAX_LENGTH
from simple_auction.models import Lot

SUPPORTED = (".xlsx", ".xlsm", ".csv")

# Leave room for the " - Excellent" condition suffix added later.
TITLE_FROM_DESC_MAX = TITLE_MAX_LENGTH - 15

# Lot field -> words that mark a column holding it (compared after lowercasing
# and dropping punctuation). Order matters: earlier fields claim columns first.
SYNONYMS: dict[str, list[str]] = {
    "lot_number": ["lot", "lot no", "lot number", "lot num", "lotnumber", "lot #"],
    "serial": ["serial", "serial no", "serial number", "sn", "s n", "serial #"],
    "book_no": ["book", "book no", "book number", "book #", "bound book"],
    "owner": ["owner", "consignor", "seller", "consigned by", "consignee"],
    "condition": ["condition", "cond", "grade"],
    "year": ["year", "yr", "year made", "manufacture year", "mfg year", "dom"],
    "make": ["make", "maker", "manufacturer", "brand", "mfg"],
    "model": ["model"],
    "title": ["title", "name", "item", "lot title", "headline"],
    "desc": ["desc", "description", "details", "notes", "lot description"],
}

FIELD_ORDER = list(SYNONYMS)

# Too general to trust inside a longer header ("Consignor Name").
_VAGUE = {"name", "item", "notes", "details", "sn", "s n", "lot", "book", "yr"}


class SheetReadError(Exception):
    """The file couldn't be read as a spreadsheet."""


@dataclass
class Sheet:
    headers: list[str]
    rows: list[list[object]] = field(default_factory=list)


def _clean(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _normalize(header: str) -> str:
    text = header.lower().replace("#", " #").replace(".", " ")
    text = re.sub(r"[^a-z0-9# ]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _raw_rows(path: Path) -> list[list[object]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as f:
            return [list(r) for r in csv.reader(f)]
    if suffix in (".xlsx", ".xlsm"):
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as e:  # openpyxl raises many kinds for bad files
            raise SheetReadError(str(e)) from e
        try:
            return [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    raise SheetReadError(f"Unsupported file type: {path.suffix or 'none'}")


def read_sheet(path: Path) -> Sheet:
    """Read the first sheet; the header row is the first of the top rows that
    names at least two known fields (else the first non-empty row)."""
    rows = [r for r in _raw_rows(path)]
    header_index = None
    for i, row in enumerate(rows[:15]):
        cells = [_clean(c) for c in row]
        if sum(1 for c in cells if c and _match_field(c)) >= 2:
            header_index = i
            break
    if header_index is None:
        header_index = next(
            (i for i, r in enumerate(rows) if any(_clean(c) for c in r)), None
        )
    if header_index is None:
        return Sheet(headers=[])
    headers = [_clean(c) for c in rows[header_index]]
    data = [r for r in rows[header_index + 1 :] if any(_clean(c) for c in r)]
    return Sheet(headers=headers, rows=data)


def _match_field(header: str) -> str | None:
    """The field a header names. An exact name beats a header ending with it
    ("Item Description"), which beats one starting with it ("Consignor
    Name"); vague words only count as the whole header."""
    norm = _normalize(header)
    if not norm:
        return None
    best: tuple[int, int, str] | None = None
    for name in FIELD_ORDER:
        for word in map(_normalize, SYNONYMS[name]):
            if norm == word:
                score = 3
            elif word in _VAGUE:
                continue
            elif norm.endswith(" " + word):
                score = 2
            elif norm.startswith(word + " "):
                score = 1
            else:
                continue
            candidate = (score, len(word), name)
            if best is None or candidate[:2] > best[:2]:
                best = candidate
    return best[2] if best else None


def guess_mapping(headers: list[str]) -> dict[str, int | None]:
    """field -> column index (or None). Each column is used at most once."""
    mapping: dict[str, int | None] = {name: None for name in FIELD_ORDER}
    for index, header in enumerate(headers):
        name = _match_field(header)
        if name and mapping[name] is None:
            mapping[name] = index
    return mapping


def _int(value: object) -> int | None:
    text = _clean(value)
    match = re.match(r"\d+", text)
    return int(match.group()) if match else None


def _year(value: object) -> int | None:
    """A 4-digit year anywhere in the text ("c. 1965", "1965-1970")."""
    match = re.search(r"\b(1[5-9]\d\d|20\d\d)\b", _clean(value))
    return int(match.group(1)) if match else None


def _title_from(desc: str) -> str:
    """A title for sheets without one: the description's first line or
    sentence, cut at a word to fit the title limit."""
    first = re.split(r"\n|(?<=\.)\s", desc.strip(), maxsplit=1)[0].strip()
    if len(first) <= TITLE_FROM_DESC_MAX:
        return first
    return first[:TITLE_FROM_DESC_MAX].rsplit(" ", 1)[0].rstrip(" ,;:-")


def lot_numbers(sheet: Sheet, column: int | None) -> list[int]:
    if column is None:
        return []
    found = (_int(r[column]) if column < len(r) else None for r in sheet.rows)
    return [n for n in found if n is not None]


def guess_auction_number(
    path: Path, numbers: list[int], step: int = AUCTION_STEP
) -> int | None:
    """From the lot numbers (rounded down to the series), else the file name."""
    if numbers:
        return (min(numbers) // step) * step or None
    digits = re.findall(r"\d+", path.stem)
    candidates = [int(d) for d in digits if int(d) >= step and int(d) % step == 0]
    return candidates[0] if candidates else None


def _keyword_condition(text: str) -> str | None:
    """Usual auction wording -> one of the default condition names."""
    t = text.lower()
    # NIB / LNIB / ANIB: (like / as) new in box.
    if "new" in t or "mint" in t or "unfired" in t or re.search(r"\b[la]?nib\b", t):
        return "Like New"
    if "excel" in t or "very good" in t:
        return "Excellent"
    if "fair" in t or "poor" in t or "worn" in t:
        return "Fair"
    if "good" in t:
        return "Good"
    return None


def parse_condition(text: str, names: list[str]) -> str:
    """Match spreadsheet condition text to one of the configured `names`.

    Blank stays blank. Exact name, then usual wording ("NIB" -> Like New) if that condition is
    configured, then a configured name inside the text ("Good+"). Anything
    else is kept as written, so no information is lost.
    """
    text = " ".join(text.split())
    by_lower = {n.lower(): n for n in names}
    if not text:
        return ""  # left for someone to pick in the app
    if text.lower() in by_lower:
        return by_lower[text.lower()]
    keyword = _keyword_condition(text)
    if keyword and keyword.lower() in by_lower:
        return by_lower[keyword.lower()]
    for name in sorted(names, key=len, reverse=True):
        if re.search(rf"(?<!\w){re.escape(name)}(?!\w)", text, re.IGNORECASE):
            return name
    return text


def lots_in_range(
    numbers: list[int], auction_no: int, step: int = AUCTION_STEP
) -> bool:
    return bool(numbers) and all(auction_no <= n < auction_no + step for n in numbers)


def build_lots(
    sheet: Sheet,
    mapping: dict[str, int | None],
    auction_no: int,
    renumber: bool,
    step: int = AUCTION_STEP,
    conditions: list[str] | None = None,
) -> list[Lot]:
    """Turn sheet rows into lots for auction `auction_no`.

    Condition text is matched to the `conditions` names (see parse_condition).

    With `renumber` (or no lot column), lots are numbered auction_no,
    auction_no + 1, ... in row order. Rows that end up with a lot number
    outside the auction's range, or a duplicate number, are skipped.
    """

    def cell(row: list[object], name: str) -> str:
        col = mapping.get(name)
        return _clean(row[col]) if col is not None and col < len(row) else ""

    lots: list[Lot] = []
    seen: set[int] = set()
    next_number = auction_no
    for row in sheet.rows:
        number = None if renumber else _int(cell(row, "lot_number"))
        if number is None:
            number = next_number
        next_number = max(next_number, number) + 1
        if not auction_no <= number < auction_no + step or number in seen:
            continue
        seen.add(number)
        lots.append(
            Lot(
                lot_number=number,
                serial=cell(row, "serial"),
                condition=parse_condition(cell(row, "condition"), conditions or []),
                title=cell(row, "title") or _title_from(cell(row, "desc")),
                desc=cell(row, "desc"),
                owner=cell(row, "owner"),
                book_no=cell(row, "book_no"),
                year=_year(cell(row, "year")),
                make=cell(row, "make"),
                model=cell(row, "model"),
            )
        )
    return lots
