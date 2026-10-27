import csv
import io
from contextlib import contextmanager
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from simple_auction.models import Lot
from simple_auction.services import storage

# The field mapping also recognizes headers used by older app workbooks.
COLUMNS = [
    ("Lot Number", "lot_number", "Lot"),
    ("Title", "title", "Title"),
    ("Description", "desc", "Desc"),
    ("Owner", "owner", "Owner"),
    ("Book Number", "book_no", "Book #"),
    ("Serial", "serial", "Serial"),
    ("Condition", "condition", "Condition"),
    ("Year", "year", "Year"),
]
HEADERS = [header for header, _field, _old in COLUMNS]


def _column_fields(headers) -> list[str | None]:
    aliases = {name: field for header, field, old in COLUMNS for name in (header, old)}
    fields = [aliases.get(header) for header in headers]
    known = [field for field in fields if field is not None]
    if "lot_number" not in known or len(known) != len(set(known)):
        raise ValueError("Workbook has missing or ambiguous lot columns")
    return fields


def _order_columns(ws) -> None:
    """Reorder app columns without discarding extra columns or cell formatting."""
    headers = [cell.value for cell in ws[1]]
    if headers[: len(HEADERS)] == HEADERS:
        return
    fields = _column_fields(headers)
    order = [
        fields.index(field) if field in fields else None
        for _header, field, _old in COLUMNS
    ] + [i for i, field in enumerate(fields) if field is None]
    dimensions = {key: copy(value) for key, value in ws.column_dimensions.items()}
    for row in ws.iter_rows():
        originals = [copy(cell) for cell in row]
        for column, source in enumerate(order, 1):
            target = ws.cell(row[0].row, column)
            if source is None:
                target.value = None
                target._style = None
                target.comment = None
                target.hyperlink = None
                continue
            cell = originals[source]
            target.value = cell.value
            target._style = copy(cell._style)
            target.comment = copy(cell.comment)
            target.hyperlink = copy(cell.hyperlink)
    ws.column_dimensions.clear()
    for column, source in enumerate(order, 1):
        old_letter = get_column_letter(source + 1) if source is not None else ""
        if old_letter in dimensions:
            letter = get_column_letter(column)
            dimension = dimensions[old_letter]
            dimension.index = letter
            dimension.min = dimension.max = column
            ws.column_dimensions[letter] = dimension
    for column, header in enumerate(HEADERS, 1):
        ws.cell(1, column, header)


class ExcelLockedError(Exception):
    """The workbook is open in Excel and can't be written."""


def auction_path(folder: Path, auction_no: int) -> Path:
    return folder / f"{auction_no}.xlsx"


def list_auctions(folder: Path) -> list[int]:
    """Auction numbers, parsed from filenames like 41000.xlsx."""
    with storage.lock:
        if not folder.exists():
            return []
        for journal in folder.glob(".*.xlsx.transaction"):
            storage.recover(journal)
        return sorted(
            int(p.stem)
            for p in folder.glob("*.xlsx")
            if p.stem.isdigit() and not p.name.startswith("~$")
        )


def create_auction(path: Path) -> None:
    with storage.lock:
        storage.recover(storage.journal_path(path))
        wb = Workbook()
        try:
            wb.active.append(HEADERS)
            _save(wb, path)
        finally:
            wb.close()


def load_auction(path: Path) -> list[Lot]:
    with storage.lock:
        storage.recover(storage.journal_path(path))
        wb = load_workbook(path, read_only=True)
        try:
            headers = next(wb.active.iter_rows(min_row=1, max_row=1, values_only=True))
            fields = _column_fields(headers)
            lot_column = fields.index("lot_number")
            rows = wb.active.iter_rows(min_row=2, values_only=True)
            return [
                _row_to_lot(row, headers) for row in rows if row[lot_column] is not None
            ]
        finally:
            wb.close()


def save_lot(path: Path, lot: Lot) -> None:
    with _edit_workbook(path) as wb:
        _put_lot(wb, lot)
        _save(wb, path)


def prepare_add(path: Path, lots: list[Lot]) -> tuple[bytes, list[Lot]]:
    with _edit_workbook(path) as wb:
        existing = {c[0].value for c in wb.active.iter_rows(min_row=2)}
        added = []
        for lot in sorted(lots, key=lambda x: x.lot_number):
            if lot.lot_number not in existing:
                wb.active.append(_lot_to_row(lot))
                existing.add(lot.lot_number)
                added.append(lot)
        return _encode(wb), added


def prepare_delete(path: Path, lot_numbers: set[int]) -> bytes:
    with _edit_workbook(path) as wb:
        if not path.exists():
            raise FileNotFoundError(path)
        rows = [
            c[0].row
            for c in wb.active.iter_rows(min_row=2)
            if c[0].value in lot_numbers
        ]
        for row in reversed(rows):
            wb.active.delete_rows(row)
        return _encode(wb)


def add_lots(path: Path, lots: list[Lot]) -> list[Lot]:
    """Append lots the workbook doesn't have yet (by lot number), in one save.
    Existing lots are left untouched. Returns the lots that were added."""
    with storage.lock:
        data, added = prepare_add(path, lots)
        _write_data(path, data)
        return added


def delete_lots(path: Path, lot_numbers: set[int]) -> None:
    with storage.lock:
        _write_data(path, prepare_delete(path, lot_numbers))


def export_auction(path: Path, destination: Path) -> None:
    """Export the active sheet as CSV without editing the source workbook."""
    with _edit_workbook(path) as wb:
        if not path.exists():
            raise FileNotFoundError(path)
        _write_data(destination, _csv_data(wb.active.values))


def export_lots(lots: list[Lot], destination: Path) -> None:
    """Export shared lots with the same CSV columns as a local auction."""
    _write_data(destination, _csv_data([HEADERS, *map(_lot_to_row, lots)]))


def _csv_data(rows) -> bytes:
    stream = io.StringIO(newline="")
    csv.writer(stream).writerows(rows)
    return stream.getvalue().encode("utf-8-sig")


@contextmanager
def _edit_workbook(path: Path):
    with storage.lock:
        storage.recover(storage.journal_path(path))
        wb = load_workbook(path) if path.exists() else Workbook()
        try:
            if not path.exists():
                wb.active.append(HEADERS)
            else:
                _order_columns(wb.active)
            yield wb
        finally:
            wb.close()


def _put_lot(wb, lot: Lot) -> None:
    ws = wb.active
    row = _lot_to_row(lot)
    for cells in ws.iter_rows(min_row=2):
        if cells[0].value == lot.lot_number:
            for cell, value in zip(cells, row):
                cell.value = value
            break
    else:
        ws.append(row)


def _encode(wb) -> bytes:
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def prepare_lot(path: Path, lot: Lot) -> bytes:
    """Encode an update without changing the original workbook."""
    with _edit_workbook(path) as wb:
        _put_lot(wb, lot)
        return _encode(wb)


def _save(wb, path: Path) -> None:
    _write_data(path, _encode(wb))


def _write_data(path: Path, data: bytes) -> None:
    try:
        storage.atomic_write(path, data)
    except PermissionError as e:
        raise ExcelLockedError(path) from e


def _lot_to_row(lot: Lot) -> list:
    return [getattr(lot, field) for _header, field, _old in COLUMNS]


def _row_to_lot(row: tuple, headers=HEADERS) -> Lot:
    values = {
        field: value
        for field, value in zip(_column_fields(headers), row)
        if field is not None
    }
    year = values.get("year")
    condition = values.get("condition")
    book_no = values.get("book_no")
    return Lot(
        lot_number=int(values["lot_number"]),
        serial=values.get("serial") or "",
        condition=str(condition).strip() if condition is not None else "",
        title=values.get("title") or "",
        desc=values.get("desc") or "",
        owner=values.get("owner") or "",
        book_no=str(book_no) if book_no is not None else "",
        year=int(year) if year else None,
    )
