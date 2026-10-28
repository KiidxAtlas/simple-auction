import csv
import io

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font

from simple_auction.models import Lot
from simple_auction.services import excel


def test_save_load_roundtrip_and_update(tmp_path):
    path = excel.auction_path(tmp_path, 41000)
    excel.save_lot(path, Lot(41000, serial="A1", condition="Good", year=1990))
    excel.save_lot(path, Lot(41001, title="Rifle", owner="Bob", book_no="7"))
    excel.save_lot(path, Lot(41000, serial="A2", condition="Fair"))

    lots = excel.load_auction(path)
    assert [x.lot_number for x in lots] == [41000, 41001]
    assert lots[0].serial == "A2"
    assert lots[0].condition == "Fair"
    assert lots[1].owner == "Bob" and lots[1].book_no == "7"


def test_delete_lots(tmp_path):
    path = excel.auction_path(tmp_path, 41000)
    for n in (41000, 41001, 41002, 41003):
        excel.save_lot(path, Lot(n))
    excel.delete_lots(path, {41001, 41002})
    assert [x.lot_number for x in excel.load_auction(path)] == [41000, 41003]


def test_list_auctions_ignores_junk(tmp_path):
    for name in ["41000.xlsx", "40000.xlsx", "notes.xlsx", "~$41000.xlsx"]:
        (tmp_path / name).touch()
    assert excel.list_auctions(tmp_path) == [40000, 41000]


EXPECTED_HEADERS = [
    "Lot Number",
    "Title",
    "Description",
    "Owner",
    "Book Number",
    "Serial",
    "Condition",
    "Year",
]
LEGACY_HEADERS = [
    "Lot",
    "Serial",
    "Condition",
    "Title",
    "Desc",
    "Owner",
    "Book #",
    "Year",
]


def test_new_workbook_column_order_matches_values(tmp_path):
    path = tmp_path / "41000.xlsx"
    excel.create_auction(path)
    lot = Lot(
        41001,
        title="Rifle",
        desc="Description text",
        owner="Owner name",
        book_no="007",
        serial="ABC",
        condition="Good",
        year=1990,
    )
    excel.save_lot(path, lot)
    with path.open("rb") as source:
        wb = load_workbook(source)
        try:
            assert list(wb.active.values) == [
                tuple(EXPECTED_HEADERS),
                (
                    41001,
                    "Rifle",
                    "Description text",
                    "Owner name",
                    "007",
                    "ABC",
                    "Good",
                    1990,
                ),
            ]
        finally:
            wb.close()
    assert excel.load_auction(path) == [lot]


@pytest.mark.parametrize("include_year", [False, True])
def test_legacy_load_export_and_save_preserve_fields_and_formatting(
    tmp_path, include_year
):
    path, exported = tmp_path / "41000.xlsx", tmp_path / "export.csv"
    wb = Workbook()
    ws = wb.active
    headers = LEGACY_HEADERS if include_year else LEGACY_HEADERS[:-1]
    values = [41001, "ABC", "Good", "Rifle", "Description text", "Owner name", 7]
    if include_year:
        values.append(1990)
    ws.append([*headers, "Staff note"])
    ws.append([*values, "Retain this extra column"])
    ws["D2"].font = Font(bold=True)
    ws["D2"].comment = Comment("Title comment", "Staff")
    ws.column_dimensions["D"].width = 35
    wb.create_sheet("Notes")["A1"] = "Retain this sheet"
    wb.save(path)
    wb.close()
    original = path.read_bytes()
    lot = Lot(
        41001,
        serial="ABC",
        condition="Good",
        title="Rifle",
        desc="Description text",
        owner="Owner name",
        book_no="7",
        year=1990 if include_year else None,
    )
    assert excel.load_auction(path) == [lot]
    excel.export_auction(path, exported)
    assert path.read_bytes() == original
    with exported.open(encoding="cp1252", newline="") as stream:
        assert list(csv.reader(stream)) == [
            EXPECTED_HEADERS[:5],
            [
                "41001",
                "Rifle",
                "Description text",
                "Owner name",
                "7",
            ],
        ]
    excel.save_lot(path, lot)
    wb = load_workbook(path)
    try:
        ws = wb.active
        assert list(next(ws.values)) == [*EXPECTED_HEADERS, "Staff note"]
        assert ws["I2"].value == "Retain this extra column"
        assert ws["B2"].font.bold
        assert ws["B2"].comment.text == "Title comment"
        assert ws.column_dimensions["B"].width == 35
        assert wb["Notes"]["A1"].value == "Retain this sheet"
    finally:
        wb.close()
    excel.add_lots(path, [Lot(41002, title="Next lot")])
    excel.delete_lots(path, {41001})
    assert excel.load_auction(path) == [Lot(41002, title="Next lot")]


def test_export_missing_source_does_not_create_workbook(tmp_path):
    destination = tmp_path / "export.csv"
    with pytest.raises(FileNotFoundError):
        excel.export_auction(tmp_path / "missing.xlsx", destination)
    assert not destination.exists()


@pytest.mark.parametrize("shared", [False, True])
@pytest.mark.parametrize("empty", [False, True])
def test_export_csv_content(tmp_path, shared, empty):
    source = tmp_path / "41000.xlsx"
    destination = tmp_path / "export.csv"
    lots = (
        []
        if empty
        else [
            Lot(
                41001,
                title='Café, "special"',
                desc="Line one\nLine two",
                owner="Zoë",
                book_no="007",
                serial="00123",
                condition="Good",
                year=1990,
            )
        ]
    )
    excel.create_auction(source)
    for lot in lots:
        excel.save_lot(source, lot)
    original = source.read_bytes()
    if shared:
        excel.export_lots(lots, destination)
    else:
        excel.export_auction(source, destination)
    expected = ",".join(EXPECTED_HEADERS[:5]) + "\r\n"
    if not empty:
        expected += '41001,"Café, ""special""",Line one Line two,Zoë,007\r\n'
    assert destination.read_bytes() == expected.encode("cp1252")
    assert source.read_bytes() == original


@pytest.mark.parametrize("shared", [False, True])
def test_export_keeps_lots_on_one_line_and_preserves_empty_cells(tmp_path, shared):
    source = tmp_path / "41000.xlsx"
    destination = tmp_path / "export.csv"
    lots = [
        Lot(
            41001,
            title='Café,\r\n"special" Rifle',
            desc="Make: Example\r\nModel: One\rCondition: Good\nWarranty: None\u2028End",
            owner="Owner\nname",
            book_no="007",
            serial="PRIVATE-SERIAL",
        ),
        Lot(41002, title="Empty description", owner="Consignor", book_no="008"),
        Lot(41003, title="Whitespace description", desc=" \t\r\n\u00a0"),
    ]
    excel.create_auction(source)
    for lot in lots:
        excel.save_lot(source, lot)
    original = source.read_bytes()

    if shared:
        excel.export_lots(lots, destination)
    else:
        excel.export_auction(source, destination)

    # Bulk Loader requires carriage returns removed, even inside quoted fields.
    lines = destination.read_text(encoding="cp1252").splitlines()
    assert [next(csv.reader([line], strict=True)) for line in lines] == [
        EXPECTED_HEADERS[:5],
        [
            "41001",
            'Café, "special" Rifle',
            "Make: Example Model: One Condition: Good Warranty: None End",
            "Owner name",
            "007",
        ],
        ["41002", "Empty description", "", "Consignor", "008"],
        ["41003", "Whitespace description", " \t \u00a0", "", ""],
    ]
    assert source.read_bytes() == original


def test_export_does_not_emit_empty_workbook_rows_as_lots(tmp_path):
    source = tmp_path / "41000.xlsx"
    destination = tmp_path / "export.csv"
    excel.save_lot(source, Lot(41001, title="First", desc="First description"))
    excel.save_lot(source, Lot(41002, title="Second", desc="Second description"))
    wb = load_workbook(source)
    wb.active.insert_rows(3)
    wb.active["A8"].font = Font(bold=True)
    wb.save(source)
    wb.close()

    excel.export_auction(source, destination)

    with destination.open(encoding="cp1252", newline="") as stream:
        assert list(csv.reader(stream)) == [
            EXPECTED_HEADERS[:5],
            ["41001", "First", "First description", "", ""],
            ["41002", "Second", "Second description", "", ""],
        ]


@pytest.mark.parametrize("shared", [False, True])
@pytest.mark.parametrize("empty", [False, True])
def test_xlsx_export_preserves_text_identifiers_and_descriptions(
    tmp_path, shared, empty
):
    source = tmp_path / "46000.xlsx"
    destination = tmp_path / "export.xlsx"
    lots = (
        []
        if empty
        else [
            Lot(
                46000,
                title='Rifle, "special"',
                desc="First line\nSecond line",
                owner="Owner",
                book_no="07212026",
                serial="PRIVATE-SERIAL",
            ),
            Lot(46001, title="Long identifier", book_no="00000000000000012345"),
            Lot(46002, title="Literal identifier", book_no="=001+002"),
            Lot(46003, title="Empty cells"),
        ]
    )
    excel.create_auction(source)
    for lot in lots:
        excel.save_lot(source, lot)
    wb = load_workbook(source)
    wb.active["I1"] = "Staff note"
    wb.active["I2"] = "Not for export"
    wb.create_sheet("Private")
    wb.save(source)
    wb.close()
    original = source.read_bytes()

    if shared:
        excel.export_lots(lots, destination)
    else:
        excel.export_auction(source, destination)

    wb = load_workbook(destination)
    try:
        assert len(wb.worksheets) == 1
        expected = [tuple(EXPECTED_HEADERS[:5])]
        if not empty:
            expected += [
                (
                    46000,
                    'Rifle, "special"',
                    "First line Second line",
                    "Owner",
                    "07212026",
                ),
                (46001, "Long identifier", None, None, "00000000000000012345"),
                (46002, "Literal identifier", None, None, "=001+002"),
                (46003, "Empty cells", None, None, None),
            ]
            for row in wb.active.iter_rows(min_row=2, max_row=4):
                assert row[4].data_type == "s"
                assert row[4].number_format == "@"
        assert list(wb.active.values) == expected
    finally:
        wb.close()
    assert source.read_bytes() == original


@pytest.mark.parametrize("shared", [False, True])
def test_xlsx_to_csv_keeps_complete_lots_on_single_lines(tmp_path, shared):
    source = tmp_path / "46000.xlsx"
    destination = tmp_path / "export.xlsx"
    description = 'Caliber: 20 Gauge\r\nType: Shotgun\rBarrel: 27"\n\nCondition: Good'
    lots = [
        Lot(
            46000,
            title='Rifle,\n"one"',
            desc=description,
            owner="Donald\u2028Hughes",
            book_no="0623-6",
        ),
        Lot(46001, title="Second rifle", desc=description, book_no="007"),
        Lot(46002, title="Empty description", owner="Owner", book_no="008"),
    ]
    excel.create_auction(source)
    for lot in lots:
        excel.save_lot(source, lot)
    original = source.read_bytes()
    if shared:
        excel.export_lots(lots, destination)
    else:
        excel.export_auction(source, destination)

    # Convert the exported cell values as Save As CSV does, without fixing them.
    converted = io.StringIO(newline="")
    wb = load_workbook(destination)
    try:
        csv.writer(converted).writerows(wb.active.values)
    finally:
        wb.close()
    lines = converted.getvalue().splitlines()
    assert [next(csv.reader([line], strict=True)) for line in lines] == [
        EXPECTED_HEADERS[:5],
        [
            "46000",
            'Rifle, "one"',
            'Caliber: 20 Gauge Type: Shotgun Barrel: 27"  Condition: Good',
            "Donald Hughes",
            "0623-6",
        ],
        [
            "46001",
            "Second rifle",
            'Caliber: 20 Gauge Type: Shotgun Barrel: 27"  Condition: Good',
            "",
            "007",
        ],
        ["46002", "Empty description", "", "Owner", "008"],
    ]
    assert source.read_bytes() == original


@pytest.mark.parametrize("shared", [False, True])
def test_proxibid_csv_has_clean_legacy_headers_and_punctuation(tmp_path, shared):
    source = tmp_path / "46000.xlsx"
    destination = tmp_path / "export.csv"
    lot = Lot(
        46000,
        title="Café’s “Special” 12‑gauge",
        desc="Maker’s finish – bright — no wear.\nSecond line.",
        owner="Zoë",
        book_no="007",
    )
    excel.save_lot(source, lot)
    original = source.read_bytes()

    if shared:
        excel.export_lots([lot], destination)
    else:
        excel.export_auction(source, destination)

    # Proxibid's legacy reader must see the real first header, not ï»¿Lot Number.
    with destination.open(encoding="cp1252", newline="") as stream:
        assert list(csv.reader(stream)) == [
            EXPECTED_HEADERS[:5],
            [
                "46000",
                'Café\'s "Special" 12-gauge',
                "Maker's finish - bright - no wear. Second line.",
                "Zoë",
                "007",
            ],
        ]
    assert source.read_bytes() == original


@pytest.mark.parametrize("shared", [False, True])
def test_unrepresentable_csv_text_preserves_existing_files(tmp_path, shared):
    source = tmp_path / "46000.xlsx"
    destination = tmp_path / "export.csv"
    lot = Lot(46000, title="東京", book_no="007")
    excel.save_lot(source, lot)
    original = source.read_bytes()
    destination.write_bytes(b"previous catalogue")

    with pytest.raises(UnicodeEncodeError):
        if shared:
            excel.export_lots([lot], destination)
        else:
            excel.export_auction(source, destination)

    assert destination.read_bytes() == b"previous catalogue"
    assert source.read_bytes() == original

    workbook = tmp_path / "export.xlsx"
    excel.export_lots([lot], workbook)
    wb = load_workbook(workbook)
    try:
        assert wb.active["B2"].value == "東京"
        assert wb.active["E2"].value == "007"
    finally:
        wb.close()


def test_export_write_failure_preserves_source(tmp_path, monkeypatch):
    source = tmp_path / "41000.xlsx"
    destination = tmp_path / "export.csv"
    excel.create_auction(source)
    original = source.read_bytes()

    def fail(*args):
        raise PermissionError("Read-only destination")

    monkeypatch.setattr(excel.storage, "atomic_write", fail)
    with pytest.raises(excel.ExcelLockedError):
        excel.export_auction(source, destination)
    assert source.read_bytes() == original
    assert not destination.exists()
